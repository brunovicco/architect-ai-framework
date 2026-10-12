from io import BytesIO

from docx import Document
from fpdf import FPDF

from .i18n import e, translate_items, translate_pattern
from .models import DecisionPlan
from .scoring import normalized_score


def _bullets(items: list[str], lang: str) -> str:
    return "\n".join(f"- {item}" for item in items) if items else f"- {e(lang, 'not_defined')}"


def _model_comparison_lines(plan: DecisionPlan, lang: str) -> list[str]:
    lines = []
    for row in plan.model_comparison:
        notes = "; ".join(f"{name}: {note}" for name, note in row.notes.items() if note) or e(lang, "not_defined")
        lines.append(f"- **{row.criterion}:** {notes}")
    return lines


def to_markdown(plan: DecisionPlan, lang: str = "en") -> str:
    c = plan.case
    r = plan.recommendation
    not_defined = e(lang, "not_defined")
    option_lines = []
    for option in plan.option_scores:
        option_lines.append(f"- **{option.name}:** {normalized_score(option.scores):.1%}")
    option_text = "\n".join(option_lines) or f"- {e(lang, 'no_option_matrix')}"
    comparison_lines = _model_comparison_lines(plan, lang)
    comparison_block = (
        f"\n## {e(lang, 'model_comparison_heading')}\n" + "\n".join(comparison_lines) + "\n"
        if comparison_lines
        else ""
    )

    return f"""# {e(lang, 'doc_title')} — {c.name}

## {e(lang, 'section_a')}
**{e(lang, 'business_problem')}:** {c.business_problem or not_defined}
**{e(lang, 'desired_outcome')}:** {c.desired_outcome or not_defined}
**{e(lang, 'users')}:** {c.users or not_defined}

## {e(lang, 'section_r')}
- {e(lang, 'quality_target')}: {c.quality_target or not_defined}
- {e(lang, 'latency_target')}: {c.latency_target or not_defined}
- {e(lang, 'cost_constraint')}: {c.cost_constraint or not_defined}
- {e(lang, 'regulatory_constraints')}: {c.regulatory_constraints or not_defined}

## {e(lang, 'section_responsibility')}
- {e(lang, 'recommended_pattern')}: **{translate_pattern(lang, r.architecture_pattern)}**

### {e(lang, 'rationale')}
{_bullets(translate_items(lang, r.rationale), lang)}

## {e(lang, 'section_i')}
- {e(lang, 'recommended_integration')}: **{translate_pattern(lang, r.integration_pattern)}**
{_bullets(translate_items(lang, r.information_strategy), lang)}

## {e(lang, 'section_model')}
{_bullets(translate_items(lang, r.model_strategy), lang)}
{comparison_block}
## {e(lang, 'section_t')}
{_bullets(translate_items(lang, r.controls), lang)}

## {e(lang, 'section_e')}
{_bullets(translate_items(lang, r.evaluations), lang)}

## {e(lang, 'section_tradeoffs')}
### {e(lang, 'option_matrix')}
{option_text}

**{e(lang, 'selected_option')}:** {plan.selected_option or e(lang, 'not_selected')}
**{e(lang, 'tradeoff_notes')}:** {plan.tradeoff_notes or not_defined}

## {e(lang, 'section_operations')}
{_bullets(translate_items(lang, r.operations), lang)}

## {e(lang, 'warnings')}
{_bullets(translate_items(lang, r.warnings), lang)}

---
{e(lang, 'footer')}
"""


def to_docx_bytes(plan: DecisionPlan, lang: str = "en") -> bytes:
    doc = Document()
    c = plan.case
    r = plan.recommendation
    not_defined = e(lang, "not_defined")

    doc.add_heading(f"{e(lang, 'doc_title')} — {c.name}", 0)

    sections = [
        (
            e(lang, "section_a"),
            [
                f"{e(lang, 'business_problem')}: {c.business_problem or not_defined}",
                f"{e(lang, 'desired_outcome')}: {c.desired_outcome or not_defined}",
                f"{e(lang, 'users')}: {c.users or not_defined}",
            ],
        ),
        (
            e(lang, "section_r"),
            [
                f"{e(lang, 'quality_target')}: {c.quality_target or not_defined}",
                f"{e(lang, 'latency_target')}: {c.latency_target or not_defined}",
                f"{e(lang, 'cost_constraint')}: {c.cost_constraint or not_defined}",
                f"{e(lang, 'regulatory_constraints')}: {c.regulatory_constraints or not_defined}",
            ],
        ),
        (
            e(lang, "section_responsibility"),
            [f"{e(lang, 'recommended_pattern')}: {translate_pattern(lang, r.architecture_pattern)}", *translate_items(lang, r.rationale)],
        ),
        (
            e(lang, "section_i"),
            [f"{e(lang, 'recommended_integration')}: {translate_pattern(lang, r.integration_pattern)}", *translate_items(lang, r.information_strategy)],
        ),
        (e(lang, "section_model"), translate_items(lang, r.model_strategy)),
        (e(lang, "section_t"), translate_items(lang, r.controls)),
        (e(lang, "section_e"), translate_items(lang, r.evaluations)),
        (e(lang, "section_operations"), translate_items(lang, r.operations)),
    ]
    for heading, items in sections:
        doc.add_heading(heading, level=1)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    comparison_lines = _model_comparison_lines(plan, lang)
    if comparison_lines:
        doc.add_heading(e(lang, "model_comparison_heading"), level=1)
        for line in comparison_lines:
            doc.add_paragraph(line.lstrip("- "), style="List Bullet")

    doc.add_heading(e(lang, "section_tradeoffs"), level=1)
    if plan.option_scores:
        for option in plan.option_scores:
            doc.add_paragraph(f"{option.name}: {normalized_score(option.scores):.1%}", style="List Bullet")
    else:
        doc.add_paragraph(e(lang, "no_option_matrix"), style="List Bullet")
    doc.add_paragraph(f"{e(lang, 'selected_option')}: {plan.selected_option or e(lang, 'not_selected')}")
    doc.add_paragraph(f"{e(lang, 'tradeoff_notes')}: {plan.tradeoff_notes or not_defined}")

    if r.warnings:
        doc.add_heading(e(lang, "warnings"), level=1)
        for item in translate_items(lang, r.warnings):
            doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph(e(lang, "footer"))

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


_PDF_TEXT_REPLACEMENTS = {
    "—": "-",
    "–": "-",
    "‘": "'",
    "’": "'",
    "“": '"',
    "”": '"',
    "…": "...",
    "•": "-",
}


def _pdf_safe(text: str) -> str:
    # fpdf2's core (non-TTF) fonts only support Latin-1. Typographic
    # characters we use elsewhere (em dash, smart quotes, bullet) are
    # normalized to ASCII first; anything still unencodable (e.g. an emoji a
    # user typed into a free-text field) falls back to "?" instead of
    # crashing PDF generation.
    for src, dst in _PDF_TEXT_REPLACEMENTS.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_write(pdf: FPDF, text: str, *, size: int = 11, bold: bool = False, gap_before: int = 2) -> None:
    pdf.set_font("Helvetica", "B" if bold else "", size)
    pdf.set_x(pdf.l_margin)  # multi_cell doesn't reset x on its own; w=0 would
    if gap_before:  # otherwise measure from the previous call's end-x and
        pdf.ln(gap_before)  # can raise "not enough horizontal space".
    pdf.multi_cell(0, 8 if size > 12 else 6, _pdf_safe(text))


def to_pdf_bytes(plan: DecisionPlan, lang: str = "en") -> bytes:
    c = plan.case
    r = plan.recommendation
    not_defined = e(lang, "not_defined")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    def heading(text: str) -> None:
        _pdf_write(pdf, text, size=14, bold=True, gap_before=6)

    def subheading(text: str) -> None:
        _pdf_write(pdf, text, size=12, bold=True, gap_before=3)

    def body(text: str) -> None:
        _pdf_write(pdf, text, size=11, bold=False, gap_before=1)

    def bullets(items: list[str]) -> None:
        for item in items or [not_defined]:
            body(f"- {item}")

    _pdf_write(pdf, f"{e(lang, 'doc_title')} - {c.name}", size=18, bold=True, gap_before=0)

    heading(e(lang, "section_a"))
    body(f"{e(lang, 'business_problem')}: {c.business_problem or not_defined}")
    body(f"{e(lang, 'desired_outcome')}: {c.desired_outcome or not_defined}")
    body(f"{e(lang, 'users')}: {c.users or not_defined}")

    heading(e(lang, "section_r"))
    body(f"{e(lang, 'quality_target')}: {c.quality_target or not_defined}")
    body(f"{e(lang, 'latency_target')}: {c.latency_target or not_defined}")
    body(f"{e(lang, 'cost_constraint')}: {c.cost_constraint or not_defined}")
    body(f"{e(lang, 'regulatory_constraints')}: {c.regulatory_constraints or not_defined}")

    heading(e(lang, "section_responsibility"))
    body(f"{e(lang, 'recommended_pattern')}: {translate_pattern(lang, r.architecture_pattern)}")
    subheading(e(lang, "rationale"))
    bullets(translate_items(lang, r.rationale))

    heading(e(lang, "section_i"))
    body(f"{e(lang, 'recommended_integration')}: {translate_pattern(lang, r.integration_pattern)}")
    bullets(translate_items(lang, r.information_strategy))

    heading(e(lang, "section_model"))
    bullets(translate_items(lang, r.model_strategy))

    comparison_lines = _model_comparison_lines(plan, lang)
    if comparison_lines:
        heading(e(lang, "model_comparison_heading"))
        for line in comparison_lines:
            body(line.lstrip("- "))

    heading(e(lang, "section_t"))
    bullets(translate_items(lang, r.controls))

    heading(e(lang, "section_e"))
    bullets(translate_items(lang, r.evaluations))

    heading(e(lang, "section_tradeoffs"))
    subheading(e(lang, "option_matrix"))
    if plan.option_scores:
        for option in plan.option_scores:
            body(f"- {option.name}: {normalized_score(option.scores):.1%}")
    else:
        body(f"- {e(lang, 'no_option_matrix')}")
    body(f"{e(lang, 'selected_option')}: {plan.selected_option or e(lang, 'not_selected')}")
    body(f"{e(lang, 'tradeoff_notes')}: {plan.tradeoff_notes or not_defined}")

    heading(e(lang, "section_operations"))
    bullets(translate_items(lang, r.operations))

    if r.warnings:
        heading(e(lang, "warnings"))
        bullets(translate_items(lang, r.warnings))

    body(e(lang, "footer"))

    return bytes(pdf.output())


def to_adr_markdown(plan: DecisionPlan, lang: str = "en") -> str:
    """Render the plan as a Nygard-style Architecture Decision Record."""
    c = plan.case
    r = plan.recommendation
    not_defined = e(lang, "not_defined")
    option_lines = []
    for option in plan.option_scores:
        option_lines.append(f"- **{option.name}:** {normalized_score(option.scores):.1%}")
    option_text = "\n".join(option_lines) or f"- {e(lang, 'no_option_matrix')}"
    decision_intro = e(lang, "adr_decision_intro").format(
        architecture=translate_pattern(lang, r.architecture_pattern),
        integration=translate_pattern(lang, r.integration_pattern),
    )

    return f"""# {e(lang, 'adr_title_prefix')}: {c.name}

**{e(lang, 'adr_status_heading')}:** {e(lang, 'adr_status_value')}

## {e(lang, 'adr_context_heading')}
- {e(lang, 'business_problem')}: {c.business_problem or not_defined}
- {e(lang, 'desired_outcome')}: {c.desired_outcome or not_defined}
- {e(lang, 'users')}: {c.users or not_defined}
- {e(lang, 'quality_target')}: {c.quality_target or not_defined}
- {e(lang, 'latency_target')}: {c.latency_target or not_defined}
- {e(lang, 'cost_constraint')}: {c.cost_constraint or not_defined}
- {e(lang, 'regulatory_constraints')}: {c.regulatory_constraints or not_defined}

## {e(lang, 'adr_decision_heading')}
{decision_intro}

{_bullets(translate_items(lang, r.rationale), lang)}

## {e(lang, 'adr_consequences_heading')}

### {e(lang, 'section_t')}
{_bullets(translate_items(lang, r.controls), lang)}

### {e(lang, 'section_e')}
{_bullets(translate_items(lang, r.evaluations), lang)}

### {e(lang, 'section_operations')}
{_bullets(translate_items(lang, r.operations), lang)}

## {e(lang, 'adr_alternatives_heading')}
{option_text}

**{e(lang, 'selected_option')}:** {plan.selected_option or e(lang, 'not_selected')}
**{e(lang, 'tradeoff_notes')}:** {plan.tradeoff_notes or not_defined}

---
{e(lang, 'footer')}
"""
